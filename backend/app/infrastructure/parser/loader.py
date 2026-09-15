"""按扩展名解析文档为纯文本。"""
from pathlib import Path

from docx import Document

from app.infrastructure.parser.pdf import extract_pdf_text


def extract_text(path: str | Path, *, preview: bool = True) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        text = extract_pdf_text(p, preview=preview)
    elif suffix in {".txt", ".md", ".markdown"}:
        text = p.read_text(encoding="utf-8", errors="ignore")
        # 生成语料可用该标记截断体积填充，避免切分出海量无意义片段
        marker = "<!-- NOTEMIND_PAD_START -->"
        if marker in text:
            text = text.split(marker, 1)[0]
        text = text.strip()
    elif suffix == ".docx":
        doc = Document(str(p))
        text = "\n".join(para.text for para in doc.paragraphs if para.text.strip()).strip()
    elif suffix == ".xlsx":
        text = _extract_xlsx(p)
    elif suffix == ".pptx":
        text = _extract_pptx(p)
    else:
        raise ValueError(f"Unsupported file type: {suffix or '(none)'}")
    # PG text / 向量入库禁止 NUL
    return (text or "").replace("\x00", "")


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"# 工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(parts).strip()


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"# 幻灯片 {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                table = shape.table
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
    return "\n".join(parts).strip()
